#!/usr/bin/env python
"""
Build the small synthetic test documents used by the test suite.

All names, identifiers and figures are fictional ("Acme Widgets Pvt Ltd",
"Director A", "Director B", "Beta Supplies LLP", "Example & Co").

Usage:
    python tests/fixtures/make_fixtures.py [OUTPUT_DIR]

Requires PyMuPDF, python-docx, openpyxl and Pillow.
"""

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

FIXTURE_DIR = Path(__file__).parent

COMPANY = "Acme Widgets Pvt Ltd"
RELATED_PARTY = "Beta Supplies LLP"
AUDITOR = "Example & Co, Chartered Accountants"

REPORT_LINES = [
    f"{COMPANY} - Annual Report FY 2024-25 (fictional sample)",
    "",
    f"{COMPANY} (CIN U00000XX0000PTC000000) manufactures widgets.",
    "Registered office: 1 Example Road, Sample City.",
    "",
    "Board of Directors:",
    "Director A (DIN 00000001) - Managing Director",
    "Director B (DIN 00000002) - Independent Director",
    "",
    f"Statutory auditor: {AUDITOR}.",
    "",
    "Related party disclosures:",
    f"{RELATED_PARTY} is a related party because Director A is a designated partner.",
    f"Purchases of raw material from {RELATED_PARTY}: INR 1,200,000.",
    "Unsecured loan received from Director B: INR 500,000.",
]


def make_pdf(path: Path):
    import fitz  # PyMuPDF

    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4
    y = 60
    for line in REPORT_LINES:
        page.insert_text((50, y), line, fontsize=11)
        y += 18
    doc.set_metadata({
        "title": "Acme Widgets annual report (synthetic)",
        "author": "Sample Author",
        "creator": "",
        "producer": "",
    })
    doc.save(str(path), garbage=4, deflate=True)
    doc.close()


def make_docx(path: Path):
    from docx import Document

    doc = Document()
    doc.add_heading("Minutes of the Board Meeting (fictional sample)", level=1)
    doc.add_paragraph(
        "Minutes of the meeting of the Board of Directors of Acme Widgets Private Limited "
        "held on 15 July 2025 at the registered office."
    )
    doc.add_paragraph("Present: Director A (Chairperson), Director B.")
    doc.add_paragraph(
        f"Resolved that the purchase contract with {RELATED_PARTY} for raw material worth "
        "INR 1,200,000 be approved. Director A, being a designated partner of "
        f"{RELATED_PARTY}, disclosed the interest and abstained from voting."
    )
    doc.add_paragraph(
        "Resolved that the unsecured loan of INR 500,000 from Director B be accepted."
    )
    doc.core_properties.author = "Sample Author"
    doc.core_properties.comments = "Synthetic sample document"
    doc.core_properties.title = "Acme Widgets board minutes (synthetic)"
    doc.save(str(path))


def make_xlsx(path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Related Parties"
    ws.append(["Related Party", "Relationship", "Nature of Transaction", "Amount (INR)"])
    ws.append([RELATED_PARTY, "Director A is designated partner", "Purchase of raw material", 1200000])
    ws.append(["Director B", "Director", "Unsecured loan received", 500000])
    ws.append(["Director A", "Managing Director", "Remuneration", 900000])
    wb.properties.creator = "Sample Author"
    wb.properties.title = "Acme Widgets related parties (synthetic)"
    wb.save(str(path))


def make_png(path: Path):
    img = Image.new("RGB", (900, 160), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 28)
    except OSError:
        font = ImageFont.load_default(size=28)
    draw.text((20, 30), "ACME WIDGETS PVT LTD", fill="black", font=font)
    draw.text((20, 90), "Director A - Managing Director", fill="black", font=font)
    img.save(str(path), optimize=True)


def build_all(out_dir: Path = FIXTURE_DIR):
    out_dir.mkdir(parents=True, exist_ok=True)
    make_pdf(out_dir / "acme_annual_report.pdf")
    make_docx(out_dir / "acme_board_minutes.docx")
    make_xlsx(out_dir / "acme_related_parties.xlsx")
    make_png(out_dir / "acme_letterhead.png")
    return out_dir


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else FIXTURE_DIR
    build_all(target)
    print(f"Fixtures written to {target}")
