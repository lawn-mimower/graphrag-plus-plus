#!/usr/bin/env python
"""
Generate the synthetic benchmark corpus and its ground truth.

Everything here is fictional: the companies, people, identifiers and amounts
were invented for this benchmark. The same people and companies appear in
several documents under different spellings ("Mr. Arjun Mehta", "Arjun K.
Mehta", "A. Mehta"; "Larkspur Castings Pvt. Ltd.", "LCPL"), and a few
different entities have deliberately similar names (Arjun Mehta / Arun Mehta,
Kestrel Alloys LLP / Kestrel Alloy Castings Limited, Larkspur Castings /
Larkspur Forgings / Larkspur Foundry Services).

Document text is written with inline markup, [[ENTITY_ID|surface form]]. The
generator renders the surface form, records every occurrence, and checks the
ground truth for consistency. Output is deterministic: running it twice gives
byte-identical files.

Usage:
    python benchmarks/generate_corpus.py            # writes benchmarks/corpus/ and benchmarks/ground_truth/
    python benchmarks/generate_corpus.py --out DIR  # writes DIR/corpus/ and DIR/ground_truth/

Outputs:
    corpus/*.pdf|*.docx|*.xlsx         the 8 documents
    ground_truth/entities.json         20 entities with type, canonical name and aliases
    ground_truth/mentions.json         66 gold mentions (one per document and entity) with the
                                       attributes an extractor could read there, the gold entity,
                                       and the relations between mentions of the same document
    ground_truth/relations.json        relation facts between entities with the documents stating them
    ground_truth/questions.json        39 questions with gold answers, evidence and the fewest
                                       documents that state the facts each one needs
    ground_truth/occurrences.json      every marked occurrence (document, page, entity, surface)
"""

import argparse
import datetime as dt
import io
import itertools
import json
import re
import sys
import textwrap
import zipfile
from pathlib import Path

BENCH_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BENCH_DIR.parent))
FIXED_TIME = dt.datetime(2025, 1, 1, 0, 0, 0)
NOTE = "Synthetic benchmark document. All names, identifiers and figures are fictional."
MARKUP = re.compile(r"\[\[(\w+)\|([^\]]+)\]\]")

# --------------------------------------------------------------------------
# Entities
# --------------------------------------------------------------------------
# aliases: surface forms that identify the entity on their own.
# ambiguous: surface forms used in the corpus that could mean more than one
# entity ("A. Mehta" is Arjun Mehta in one document and Arun Mehta in
# another). They are recorded as occurrences but never used to score answers.
ENTITIES = [
    {"id": "P1", "type": "Person", "name": "Arjun Mehta",
     "aliases": ["Arjun Mehta", "Mr. Arjun Mehta", "Arjun K. Mehta"],
     "ambiguous": ["A. Mehta", "Mr. Mehta"]},
    {"id": "P2", "type": "Person", "name": "Meera Iyer",
     "aliases": ["Meera Iyer", "Ms. Meera Iyer", "M. Iyer", "Ms. Iyer"]},
    {"id": "P3", "type": "Person", "name": "Rohan Pillai",
     "aliases": ["Rohan Pillai", "Mr. Rohan Pillai", "R. Pillai", "Rohan K. Pillai"]},
    {"id": "P4", "type": "Person", "name": "Kavita Bose",
     "aliases": ["Kavita Bose", "Ms. Kavita Bose", "K. Bose", "Ms. K. Bose"]},
    {"id": "P5", "type": "Person", "name": "Sameer Joshi",
     "aliases": ["Sameer Joshi", "Mr. Sameer Joshi", "S. Joshi"]},
    {"id": "P6", "type": "Person", "name": "Arun Mehta",
     "aliases": ["Arun Mehta", "Mr. Arun Mehta"],
     "ambiguous": ["A. Mehta"]},
    {"id": "P7", "type": "Person", "name": "Leela Narayan",
     "aliases": ["Leela Narayan", "Ms. Leela Narayan", "L. Narayan", "Ms L. Narayan"]},
    {"id": "P8", "type": "Person", "name": "Tomas Fernandes",
     "aliases": ["Tomas Fernandes", "Mr. Tomas Fernandes", "T. Fernandes", "Mr. Fernandes"]},
    {"id": "P9", "type": "Person", "name": "Vikram Sethi",
     "aliases": ["Vikram Sethi", "Mr Vikram Sethi", "V. Sethi"]},
    {"id": "C1", "type": "Company", "name": "Larkspur Castings Private Limited",
     "aliases": ["Larkspur Castings Private Limited", "Larkspur Castings Pvt. Ltd.",
                 "Larkspur Castings Pvt Ltd", "Larkspur Castings Private Ltd.",
                 "LARKSPUR CASTINGS PRIVATE LIMITED", "Larkspur Castings", "LCPL"]},
    {"id": "C2", "type": "Company", "name": "Tern Capital Private Limited",
     "aliases": ["Tern Capital Private Limited", "Tern Capital Pvt Ltd", "TERN CAPITAL PVT LTD",
                 "Tern Capital", "TCPL"]},
    {"id": "C3", "type": "Company", "name": "Harbourline Logistics Private Limited",
     "aliases": ["Harbourline Logistics Private Limited", "Harbourline Logistics Pvt Ltd",
                 "Harbourline Logistics Pvt. Ltd.", "M/s Harbourline Logistics Pvt. Ltd.",
                 "Harbourline Logistics", "Harbourline", "HLPL"]},
    {"id": "C4", "type": "Company", "name": "Larkspur Foundry Services Private Limited",
     "aliases": ["Larkspur Foundry Services Private Limited", "Larkspur Foundry Services Pvt Ltd",
                 "Larkspur Foundry Services", "LFSPL"]},
    {"id": "C5", "type": "Company", "name": "Kestrel Alloys LLP",
     "aliases": ["Kestrel Alloys LLP", "M/s Kestrel Alloys LLP", "Kestrel Alloys"]},
    {"id": "C6", "type": "Company", "name": "Bluefin Traders Private Limited",
     "aliases": ["Bluefin Traders Private Limited", "Bluefin Traders Pvt Ltd", "Bluefin Traders"]},
    {"id": "C7", "type": "Organization", "name": "Sen & Varma LLP",
     "aliases": ["Sen & Varma LLP", "SEN & VARMA LLP", "Sen and Varma LLP", "Sen & Varma"]},
    {"id": "C8", "type": "Organization", "name": "Patel Rao & Associates",
     "aliases": ["Patel Rao & Associates", "Patel Rao and Associates", "Patel Rao"]},
    {"id": "C9", "type": "Company", "name": "Kestrel Alloy Castings Limited",
     "aliases": ["Kestrel Alloy Castings Limited", "Kestrel Alloy Castings Ltd",
                 "Kestrel Alloy Castings"]},
    {"id": "C10", "type": "Company", "name": "Larkspur Forgings Limited",
     "aliases": ["Larkspur Forgings Limited", "Larkspur Forgings Ltd", "Larkspur Forgings"]},
    {"id": "C11", "type": "Organization", "name": "Seabright Cooperative Bank Limited",
     "aliases": ["Seabright Cooperative Bank Limited", "Seabright Co-operative Bank Ltd",
                 "SEABRIGHT COOPERATIVE BANK LIMITED", "Seabright Cooperative Bank",
                 "Seabright Bank"]},
]

# --------------------------------------------------------------------------
# Documents
# --------------------------------------------------------------------------
# PDF/DOCX pages are lists of ("h" | "p", text). XLSX pages are sheets.
# "mentions": the gold mention for each entity in the document: the name an
# extractor would most likely report and the attributes stated there.
# "relations": (from, type, to, attributes) stated in the document.
DOCUMENTS = [
    {
        "name": "larkspur_annual_report_fy2025", "format": "pdf",
        "title": "Larkspur Castings annual report FY 2024-25 (synthetic)",
        "pages": [
            [
                ("h", "[[C1|Larkspur Castings Private Limited]]"),
                ("p", "Annual Report FY 2024-25 (extract)"),
                ("p", NOTE),
                ("h", "1. Company overview"),
                ("p", "[[C1|Larkspur Castings Private Limited]] (\"[[C1|LCPL]]\" or \"the Company\") makes "
                      "grey-iron and alloy castings for pump and valve manufacturers. Registered office: "
                      "14 Canal Road, Riverton."),
                ("h", "2. Holding company and subsidiaries"),
                ("p", "[[C2|Tern Capital Private Limited]] holds 62% of the equity share capital of the "
                      "Company and is its holding company."),
                ("p", "The Company has two subsidiaries: [[C3|Harbourline Logistics Private Limited]] "
                      "(wholly owned, 100%) and [[C4|Larkspur Foundry Services Private Limited]] "
                      "(76% held by the Company)."),
                ("h", "3. Board of Directors"),
                ("p", "[[P1|Mr. Arjun Mehta]] (DIN 07410101) - Managing Director"),
                ("p", "[[P2|Ms. Meera Iyer]] (DIN 08220202) - Independent Director; Chairperson of the "
                      "Audit Committee"),
                ("p", "[[P3|Mr. Rohan Pillai]] (DIN 06530303) - Whole-time Director and Chief Financial "
                      "Officer"),
                ("p", "[[P7|Ms. Leela Narayan]] (DIN 09140404) - Nominee Director representing "
                      "[[C2|Tern Capital Private Limited]]"),
                ("h", "4. Key managerial personnel"),
                ("p", "[[P5|Mr. Sameer Joshi]] - Company Secretary"),
            ],
            [
                ("h", "5. Customers"),
                ("p", "The largest customer during the year was [[C9|Kestrel Alloy Castings Limited]], "
                      "an unrelated party, which accounted for 18% of revenue."),
                ("h", "6. Statutory auditors"),
                ("p", "[[C7|Sen & Varma LLP]], Chartered Accountants, were re-appointed as statutory "
                      "auditors of the Company. The independent auditor's report for FY 2024-25 was "
                      "signed by [[P6|Mr. Arun Mehta]], Partner (Membership No. 104455)."),
                ("h", "7. Related party transactions (summary)"),
                ("p", "The Company purchased alloy ingots worth INR 12,600,000 from "
                      "[[C5|Kestrel Alloys LLP]], in which [[P1|Mr. Arjun Mehta]] is a designated "
                      "partner. The Company gave a loan of INR 5,000,000 to its subsidiary [[C4|LFSPL]] "
                      "and a corporate guarantee for the term loan of its subsidiary [[C3|HLPL]]. Full "
                      "details are in the related party schedule."),
                ("p", "For and on behalf of the Board of Directors"),
                ("p", "[[P1|Arjun Mehta]], Managing Director; [[P3|Rohan Pillai]], Whole-time Director "
                      "and CFO. Riverton, 30 May 2025."),
            ],
        ],
        "mentions": [
            ("C1", "Larkspur Castings Private Limited", {"registered_office": "14 Canal Road, Riverton"}),
            ("C2", "Tern Capital Private Limited", {"role": "Holding company", "holding": "62%"}),
            ("C3", "Harbourline Logistics Private Limited", {"role": "Wholly owned subsidiary"}),
            ("C4", "Larkspur Foundry Services Private Limited", {"role": "Subsidiary", "holding": "76%"}),
            ("P1", "Mr. Arjun Mehta", {"din": "07410101", "role": "Managing Director"}),
            ("P2", "Ms. Meera Iyer", {"din": "08220202", "role": "Independent Director"}),
            ("P3", "Mr. Rohan Pillai", {"din": "06530303", "role": "Whole-time Director and CFO"}),
            ("P7", "Ms. Leela Narayan", {"din": "09140404", "role": "Nominee Director"}),
            ("P5", "Mr. Sameer Joshi", {"role": "Company Secretary"}),
            ("C9", "Kestrel Alloy Castings Limited", {"role": "Customer"}),
            ("C7", "Sen & Varma LLP", {"role": "Statutory auditor"}),
            ("P6", "Mr. Arun Mehta", {"role": "Partner", "membership_no": "104455"}),
            ("C5", "Kestrel Alloys LLP", {"role": "Related party"}),
        ],
        "relations": [
            ("C2", "HOLDING_COMPANY_OF", "C1", {"holding": "62%"}),
            ("C1", "PARENT_OF", "C3", {"holding": "100%"}),
            ("C1", "PARENT_OF", "C4", {"holding": "76%"}),
            ("P1", "MANAGING_DIRECTOR_OF", "C1", {}),
            ("P2", "INDEPENDENT_DIRECTOR_OF", "C1", {}),
            ("P3", "WHOLE_TIME_DIRECTOR_OF", "C1", {}),
            ("P3", "CFO_OF", "C1", {}),
            ("P7", "NOMINEE_DIRECTOR_OF", "C1", {}),
            ("P7", "NOMINEE_OF", "C2", {}),
            ("P5", "COMPANY_SECRETARY_OF", "C1", {}),
            ("C9", "CUSTOMER_OF", "C1", {}),
            ("C7", "AUDITOR_OF", "C1", {}),
            ("P6", "PARTNER_OF", "C7", {}),
            ("C1", "PURCHASED_FROM", "C5", {"amount_inr": 12600000}),
            ("P1", "DESIGNATED_PARTNER_OF", "C5", {}),
            ("C1", "LOAN_TO", "C4", {"amount_inr": 5000000}),
            ("C1", "GUARANTOR_FOR", "C3", {}),
        ],
    },
    {
        "name": "larkspur_board_minutes_jul2025", "format": "docx",
        "title": "Larkspur Castings board minutes, July 2025 (synthetic)",
        "pages": [[
            ("h", "Minutes of the Board Meeting"),
            ("p", NOTE),
            ("p", "Minutes of the meeting of the Board of Directors of [[C1|Larkspur Castings Pvt. Ltd.]] "
                  "held on 18 July 2025 at 14 Canal Road, Riverton."),
            ("p", "Present: [[P1|Arjun Mehta]] (Chairperson), [[P2|M. Iyer]], [[P3|R. Pillai]], "
                  "[[P7|L. Narayan]]. In attendance: [[P5|S. Joshi]], Company Secretary."),
            ("p", "1. Related party contract. The Board approved the renewal of the supply contract with "
                  "[[C5|M/s Kestrel Alloys LLP]] for alloy ingots up to INR 15,000,000 for FY 2025-26. "
                  "[[P1|Mr. Mehta]] disclosed his interest as a designated partner of "
                  "[[C5|Kestrel Alloys]] and did not take part in the vote."),
            ("p", "2. Corporate guarantee. The Board approved a corporate guarantee of INR 40,000,000 in "
                  "favour of [[C11|Seabright Cooperative Bank Limited]] for the term loan sanctioned to "
                  "the wholly owned subsidiary [[C3|Harbourline Logistics Pvt Ltd]]."),
            ("p", "3. Customers. The Board reviewed receivables from [[C9|Kestrel Alloy Castings Ltd]], "
                  "the largest customer, and noted the sale of foundry scrap worth INR 850,000 to "
                  "[[C10|Larkspur Forgings Limited]]. Neither company is related to [[C1|LCPL]]; "
                  "[[C10|Larkspur Forgings Limited]] only has a similar name."),
            ("p", "4. Audit committee. [[P2|Ms. Iyer]] reported that the Audit Committee had reviewed the "
                  "quarterly results with the statutory auditors, [[C7|Sen and Varma LLP]]."),
            ("p", "The meeting ended with a vote of thanks to the Chair."),
        ]],
        "mentions": [
            ("C1", "Larkspur Castings Pvt. Ltd.", {}),
            ("P1", "Arjun Mehta", {"role": "Chairperson"}),
            ("P2", "M. Iyer", {"role": "Director"}),
            ("P3", "R. Pillai", {"role": "Director"}),
            ("P7", "L. Narayan", {"role": "Director"}),
            ("P5", "S. Joshi", {"role": "Company Secretary"}),
            ("C5", "M/s Kestrel Alloys LLP", {"role": "Supplier"}),
            ("C11", "Seabright Cooperative Bank Limited", {"role": "Lender"}),
            ("C3", "Harbourline Logistics Pvt Ltd", {"role": "Wholly owned subsidiary"}),
            ("C9", "Kestrel Alloy Castings Ltd", {"role": "Customer"}),
            ("C10", "Larkspur Forgings Limited", {"role": "Customer"}),
            ("C7", "Sen and Varma LLP", {"role": "Statutory auditor"}),
        ],
        "relations": [
            ("P1", "CHAIRPERSON_OF", "C1", {}),
            ("P2", "DIRECTOR_OF", "C1", {}),
            ("P3", "DIRECTOR_OF", "C1", {}),
            ("P7", "DIRECTOR_OF", "C1", {}),
            ("P5", "COMPANY_SECRETARY_OF", "C1", {}),
            ("C1", "CONTRACT_WITH", "C5", {"amount_inr": 15000000}),
            ("P1", "DESIGNATED_PARTNER_OF", "C5", {}),
            ("C1", "GUARANTOR_FOR", "C3", {"amount_inr": 40000000}),
            ("C11", "LENDER_TO", "C3", {}),
            ("C1", "PARENT_OF", "C3", {"holding": "100%"}),
            ("C9", "CUSTOMER_OF", "C1", {}),
            ("C10", "CUSTOMER_OF", "C1", {"amount_inr": 850000}),
            ("C7", "AUDITOR_OF", "C1", {}),
        ],
    },
    {
        "name": "larkspur_related_parties_fy2025", "format": "xlsx",
        "title": "Larkspur Castings related party schedule FY 2024-25 (synthetic)",
        "pages": [
            {"sheet": "Related Parties FY2025", "rows": [
                ["[[C1|Larkspur Castings Private Limited]] - Related party transactions FY 2024-25",
                 None, None, None],
                ["Related party", "Relationship", "Nature of transaction", "Amount (INR)"],
                ["[[C2|Tern Capital Pvt Ltd]]", "Holding company (62%)", "Dividend paid", 3100000],
                ["[[C3|Harbourline Logistics Pvt. Ltd.]]", "Wholly owned subsidiary",
                 "Corporate guarantee given to [[C11|Seabright Cooperative Bank]]", 40000000],
                ["[[C4|Larkspur Foundry Services Pvt Ltd]]", "Subsidiary (76%)", "Loan given", 5000000],
                ["[[C5|Kestrel Alloys LLP]]", "Firm in which a director ([[P1|A. Mehta]]) is a designated partner",
                 "Purchase of alloy ingots", 12600000],
                ["[[P1|Arjun Mehta]]", "Managing Director", "Remuneration", 9600000],
                ["[[P3|Rohan Pillai]]", "Whole-time Director and CFO", "Remuneration", 6200000],
                ["[[P2|Meera Iyer]]", "Independent Director", "Sitting fees", 480000],
                ["[[P5|Sameer Joshi]]", "Company Secretary", "Remuneration", 2400000],
                ["Note: [[C9|Kestrel Alloy Castings Ltd]] and [[C10|Larkspur Forgings Ltd]] are customers "
                 "and are not related parties.", None, None, None],
                [NOTE, None, None, None],
            ]},
            {"sheet": "LLP partners", "rows": [
                ["LLP", "Designated partner"],
                ["[[C5|Kestrel Alloys LLP]]", "[[P1|Arjun Mehta]]"],
                ["[[C5|Kestrel Alloys LLP]]", "[[P8|Tomas Fernandes]]"],
            ]},
        ],
        "mentions": [
            ("C1", "Larkspur Castings Private Limited", {}),
            ("C2", "Tern Capital Pvt Ltd", {"role": "Holding company", "holding": "62%"}),
            ("C3", "Harbourline Logistics Pvt. Ltd.", {"role": "Wholly owned subsidiary"}),
            ("C11", "Seabright Cooperative Bank", {"role": "Lender"}),
            ("C4", "Larkspur Foundry Services Pvt Ltd", {"role": "Subsidiary", "holding": "76%"}),
            ("C5", "Kestrel Alloys LLP", {"role": "Related party"}),
            ("P1", "Arjun Mehta", {"role": "Managing Director"}),
            ("P3", "Rohan Pillai", {"role": "Whole-time Director and CFO"}),
            ("P2", "Meera Iyer", {"role": "Independent Director"}),
            ("P5", "Sameer Joshi", {"role": "Company Secretary"}),
            ("C9", "Kestrel Alloy Castings Ltd", {"role": "Customer"}),
            ("C10", "Larkspur Forgings Ltd", {"role": "Customer"}),
            ("P8", "Tomas Fernandes", {"role": "Designated partner"}),
        ],
        "relations": [
            ("C2", "HOLDING_COMPANY_OF", "C1", {"holding": "62%"}),
            ("C1", "DIVIDEND_PAID_TO", "C2", {"amount_inr": 3100000}),
            ("C1", "GUARANTOR_FOR", "C3", {"amount_inr": 40000000}),
            ("C1", "PARENT_OF", "C3", {"holding": "100%"}),
            ("C11", "LENDER_TO", "C3", {}),
            ("C1", "LOAN_TO", "C4", {"amount_inr": 5000000}),
            ("C1", "PARENT_OF", "C4", {"holding": "76%"}),
            ("C1", "PURCHASED_FROM", "C5", {"amount_inr": 12600000}),
            ("P1", "DESIGNATED_PARTNER_OF", "C5", {}),
            ("P8", "DESIGNATED_PARTNER_OF", "C5", {}),
            ("P1", "MANAGING_DIRECTOR_OF", "C1", {"remuneration_inr": 9600000}),
            ("P3", "WHOLE_TIME_DIRECTOR_OF", "C1", {"remuneration_inr": 6200000}),
            ("P3", "CFO_OF", "C1", {}),
            ("P2", "INDEPENDENT_DIRECTOR_OF", "C1", {"sitting_fees_inr": 480000}),
            ("P5", "COMPANY_SECRETARY_OF", "C1", {"remuneration_inr": 2400000}),
            ("C9", "CUSTOMER_OF", "C1", {}),
            ("C10", "CUSTOMER_OF", "C1", {}),
        ],
    },
    {
        "name": "harbourline_annual_report_fy2025", "format": "pdf",
        "title": "Harbourline Logistics annual report FY 2024-25 (synthetic)",
        "pages": [[
            ("h", "[[C3|Harbourline Logistics Private Limited]]"),
            ("p", "Annual Report FY 2024-25 (extract)"),
            ("p", NOTE),
            ("p", "[[C3|Harbourline Logistics Private Limited]] (\"[[C3|HLPL]]\") provides inland freight "
                  "services. It is a wholly owned subsidiary of [[C1|Larkspur Castings Pvt Ltd]]."),
            ("h", "Board of Directors"),
            ("p", "[[P1|Arjun K. Mehta]] (DIN 07410101) - Director"),
            ("p", "[[P4|Kavita Bose]] (DIN 08850505) - Director"),
            ("h", "Statutory auditors"),
            ("p", "[[C8|Patel Rao & Associates]], Chartered Accountants."),
            ("h", "Borrowings"),
            ("p", "[[C11|Seabright Co-operative Bank Ltd]] sanctioned a term loan of INR 40,000,000 to "
                  "[[C3|HLPL]], secured by a corporate guarantee from the holding company, [[C1|LCPL]]."),
            ("h", "Related party transactions"),
            ("p", "Freight brokerage services received from [[C6|Bluefin Traders Private Limited]], a "
                  "company in which [[P4|Ms. Kavita Bose]] is a director: INR 1,850,000."),
            ("p", "Corporate guarantee received from the holding company, "
                  "[[C1|Larkspur Castings Pvt Ltd]]: INR 40,000,000."),
        ]],
        "mentions": [
            ("C3", "Harbourline Logistics Private Limited", {}),
            ("C1", "Larkspur Castings Pvt Ltd", {"role": "Holding company"}),
            ("P1", "Arjun K. Mehta", {"din": "07410101", "role": "Director"}),
            ("P4", "Kavita Bose", {"din": "08850505", "role": "Director"}),
            ("C8", "Patel Rao & Associates", {"role": "Statutory auditor"}),
            ("C11", "Seabright Co-operative Bank Ltd", {"role": "Lender"}),
            ("C6", "Bluefin Traders Private Limited", {"role": "Related party"}),
        ],
        "relations": [
            ("C1", "PARENT_OF", "C3", {"holding": "100%"}),
            ("C1", "GUARANTOR_FOR", "C3", {"amount_inr": 40000000}),
            ("P1", "DIRECTOR_OF", "C3", {}),
            ("P4", "DIRECTOR_OF", "C3", {}),
            ("C8", "AUDITOR_OF", "C3", {}),
            ("C11", "LENDER_TO", "C3", {"amount_inr": 40000000}),
            ("C3", "PURCHASED_SERVICES_FROM", "C6", {"amount_inr": 1850000}),
            ("P4", "DIRECTOR_OF", "C6", {}),
        ],
    },
    {
        "name": "sen_varma_engagement_letter", "format": "docx",
        "title": "Audit engagement letter (synthetic)",
        "pages": [[
            ("h", "[[C7|SEN & VARMA LLP]]"),
            ("p", "Chartered Accountants, 7 Mill Lane, Riverton"),
            ("p", NOTE),
            ("p", "To the Boards of Directors of [[C1|Larkspur Castings Private Limited]] and "
                  "[[C4|Larkspur Foundry Services Private Limited]]"),
            ("p", "Subject: Statutory audit engagement for FY 2025-26"),
            ("p", "We confirm our re-appointment as statutory auditors of both companies. "
                  "[[P6|Arun Mehta]], Partner, will lead both engagements."),
            ("p", "Proposed fees: [[C1|Larkspur Castings]] INR 1,450,000; [[C4|LFSPL]] INR 380,000, plus "
                  "applicable taxes."),
            ("p", "Accepted on behalf of [[C1|Larkspur Castings Private Limited]] by "
                  "[[P3|Rohan Pillai]], Chief Financial Officer."),
            ("p", "Yours faithfully, for [[C7|Sen & Varma LLP]]"),
            ("p", "[[P6|A. Mehta]], Partner"),
        ]],
        "mentions": [
            ("C7", "SEN & VARMA LLP", {"role": "Statutory auditor",
                                       "address": "7 Mill Lane, Riverton"}),
            ("C1", "Larkspur Castings Private Limited", {}),
            ("C4", "Larkspur Foundry Services Private Limited", {}),
            ("P6", "Arun Mehta", {"role": "Partner"}),
            ("P3", "Rohan Pillai", {"role": "Chief Financial Officer"}),
        ],
        "relations": [
            ("C7", "AUDITOR_OF", "C1", {"fee_inr": 1450000}),
            ("C7", "AUDITOR_OF", "C4", {"fee_inr": 380000}),
            ("P6", "PARTNER_OF", "C7", {}),
            ("P3", "CFO_OF", "C1", {}),
        ],
    },
    {
        "name": "larkspur_shareholding_mar2025", "format": "xlsx",
        "title": "Larkspur Castings shareholding pattern (synthetic)",
        "pages": [
            {"sheet": "Shareholding 31-Mar-2025", "rows": [
                ["[[C1|LARKSPUR CASTINGS PRIVATE LIMITED]] - Shareholding pattern as at 31 March 2025",
                 None, None, None],
                ["Shareholder", "Category", "Shares held", "% holding"],
                ["[[C2|TERN CAPITAL PVT LTD]]", "Promoter (holding company)", 6200000, "62%"],
                ["[[P1|Arjun K. Mehta]]", "Promoter group", 1500000, "15%"],
                ["Public and others", "Public", 2300000, "23%"],
                ["Total", None, 10000000, "100%"],
                [NOTE, None, None, None],
            ]},
            {"sheet": "Holding company board", "rows": [
                ["Company", "Director", "Designation"],
                ["[[C2|Tern Capital Private Limited]]", "[[P7|Ms L. Narayan]]", "Director"],
                ["[[C2|Tern Capital Private Limited]]", "[[P9|Mr Vikram Sethi]]", "Managing Director"],
            ]},
        ],
        "mentions": [
            ("C1", "LARKSPUR CASTINGS PRIVATE LIMITED", {}),
            ("C2", "TERN CAPITAL PVT LTD", {"role": "Promoter (holding company)", "holding": "62%"}),
            ("P1", "Arjun K. Mehta", {"role": "Promoter group", "holding": "15%"}),
            ("P7", "Ms L. Narayan", {"role": "Director"}),
            ("P9", "Mr Vikram Sethi", {"role": "Managing Director"}),
        ],
        "relations": [
            ("C2", "HOLDING_COMPANY_OF", "C1", {"holding": "62%"}),
            ("P1", "SHAREHOLDER_OF", "C1", {"holding": "15%"}),
            ("P7", "DIRECTOR_OF", "C2", {}),
            ("P9", "MANAGING_DIRECTOR_OF", "C2", {}),
        ],
    },
    {
        "name": "lfspl_board_minutes_may2025", "format": "docx",
        "title": "Larkspur Foundry Services board minutes, May 2025 (synthetic)",
        "pages": [[
            ("h", "Minutes of the Board Meeting"),
            ("p", NOTE),
            ("p", "Minutes of the meeting of the Board of Directors of "
                  "[[C4|Larkspur Foundry Services Pvt Ltd]] (\"[[C4|LFSPL]]\") held on 22 May 2025."),
            ("p", "Present: [[P3|Rohan K. Pillai]] (Chairperson) and [[P8|Tomas Fernandes]]."),
            ("p", "1. Holding company. The Board noted that [[C1|Larkspur Castings Pvt Ltd]] holds 76% of "
                  "the equity shares of the Company and has given it a loan of INR 5,000,000 at 9% "
                  "interest."),
            ("p", "2. Supply contract. [[P8|Mr. Fernandes]] disclosed that he is a designated partner of "
                  "[[C5|Kestrel Alloys]], from which the Company buys alloy scrap worth about "
                  "INR 2,100,000 a year. The Board approved the contract; [[P8|Mr. Fernandes]] did not "
                  "vote."),
            ("p", "3. Auditors. The Board noted that [[C7|Sen & Varma]] will continue as statutory "
                  "auditors for FY 2025-26."),
        ]],
        "mentions": [
            ("C4", "Larkspur Foundry Services Pvt Ltd", {}),
            ("P3", "Rohan K. Pillai", {"role": "Chairperson"}),
            ("P8", "Tomas Fernandes", {"role": "Director"}),
            ("C1", "Larkspur Castings Pvt Ltd", {"role": "Holding company", "holding": "76%"}),
            ("C5", "Kestrel Alloys", {"role": "Supplier"}),
            ("C7", "Sen & Varma", {"role": "Statutory auditor"}),
        ],
        "relations": [
            ("P3", "CHAIRPERSON_OF", "C4", {}),
            ("P8", "DIRECTOR_OF", "C4", {}),
            ("C1", "PARENT_OF", "C4", {"holding": "76%"}),
            ("C1", "LOAN_TO", "C4", {"amount_inr": 5000000}),
            ("P8", "DESIGNATED_PARTNER_OF", "C5", {}),
            ("C4", "PURCHASED_FROM", "C5", {"amount_inr": 2100000}),
            ("C7", "AUDITOR_OF", "C4", {}),
        ],
    },
    {
        "name": "seabright_sanction_letter", "format": "pdf",
        "title": "Term loan sanction letter (synthetic)",
        "pages": [[
            ("h", "[[C11|SEABRIGHT COOPERATIVE BANK LIMITED]]"),
            ("p", "Riverton Main Branch"),
            ("p", NOTE),
            ("h", "Sanction letter - Term loan"),
            ("p", "To: [[C3|M/s Harbourline Logistics Pvt. Ltd.]], attention: [[P4|Ms. K. Bose]], Director"),
            ("p", "We are pleased to sanction a term loan of INR 40,000,000 for a tenor of 60 months to "
                  "purchase delivery vehicles."),
            ("p", "Security: corporate guarantee of [[C1|Larkspur Castings Private Ltd.]] (holding company "
                  "of the borrower)."),
            ("p", "Covenant: the borrower's freight brokerage arrangement with "
                  "[[C6|Bluefin Traders Pvt Ltd]], a related party, must not exceed INR 2,500,000 a year."),
            ("p", "For [[C11|Seabright Bank]], Branch Manager"),
        ]],
        "mentions": [
            ("C11", "SEABRIGHT COOPERATIVE BANK LIMITED", {"role": "Lender"}),
            ("C3", "M/s Harbourline Logistics Pvt. Ltd.", {"role": "Borrower"}),
            ("P4", "Ms. K. Bose", {"role": "Director"}),
            ("C1", "Larkspur Castings Private Ltd.", {"role": "Guarantor"}),
            ("C6", "Bluefin Traders Pvt Ltd", {"role": "Related party"}),
        ],
        "relations": [
            ("C11", "LENDER_TO", "C3", {"amount_inr": 40000000}),
            ("P4", "DIRECTOR_OF", "C3", {}),
            ("C1", "GUARANTOR_FOR", "C3", {}),
            ("C1", "PARENT_OF", "C3", {}),
            ("C3", "PURCHASED_SERVICES_FROM", "C6", {"limit_inr": 2500000}),
        ],
    },
]

# --------------------------------------------------------------------------
# Questions
# --------------------------------------------------------------------------
# answer_type: entity (one gold entity), entity_any (any one of several),
# entity_list (set, scored by F1), number, string (exact token), yes_no.
# evidence: entities the context must contain to answer.
# paths: alternative lists of (entity, entity) pairs that a retrieved
# subgraph must connect (any order) for a structural hit.
QUESTIONS = [
    # lookup: one relation, usually one document
    ("Q01", "lookup", "Who is the Managing Director of Larkspur Castings Private Limited?",
     "entity", ["P1"], ["P1", "C1"], [[("P1", "C1")]]),
    ("Q02", "lookup", "Who is the Company Secretary of Larkspur Castings?",
     "entity", ["P5"], ["P5", "C1"], [[("P5", "C1")]]),
    ("Q03", "lookup", "Which firm is the statutory auditor of Harbourline Logistics?",
     "entity", ["C8"], ["C8", "C3"], [[("C8", "C3")]]),
    ("Q04", "lookup", "Which company is the holding company of Larkspur Castings?",
     "entity", ["C2"], ["C2", "C1"], [[("C2", "C1")]]),
    ("Q05", "lookup", "Who is the Chairperson of the Audit Committee of Larkspur Castings?",
     "entity", ["P2"], ["P2", "C1"], [[("P2", "C1")]]),
    ("Q06", "lookup", "Which partner of the audit firm signed the auditor's report of Larkspur Castings "
                      "for FY 2024-25?",
     "entity", ["P6"], ["P6", "C7"], [[("P6", "C7")]]),
    ("Q07", "lookup", "Which bank sanctioned a term loan to Harbourline Logistics?",
     "entity", ["C11"], ["C11", "C3"], [[("C11", "C3")]]),
    ("Q08", "lookup", "Who chaired the board meeting of Larkspur Foundry Services held on 22 May 2025?",
     "entity", ["P3"], ["P3", "C4"], [[("P3", "C4")]]),
    # attribute: a value stated next to an entity
    ("Q09", "attribute", "What is the DIN of Meera Iyer?",
     "string", "08220202", ["P2"], []),
    ("Q10", "attribute", "What percentage of the equity of Larkspur Castings is held by Tern Capital?",
     "number", 62, ["C2", "C1"], [[("C2", "C1")]]),
    ("Q11", "attribute", "What was the value, in INR, of the alloy ingots Larkspur Castings bought from "
                         "Kestrel Alloys LLP in FY 2024-25?",
     "number", 12600000, ["C1", "C5"], [[("C1", "C5")]]),
    ("Q12", "attribute", "What is the amount, in INR, of the term loan sanctioned to Harbourline Logistics?",
     "number", 40000000, ["C11", "C3"], [[("C11", "C3")]]),
    ("Q13", "attribute", "How much remuneration, in INR, did Larkspur Castings pay its Managing Director?",
     "number", 9600000, ["P1", "C1"], [[("P1", "C1")]]),
    ("Q14", "attribute", "How much did Harbourline Logistics pay Bluefin Traders for freight brokerage, "
                         "in INR?",
     "number", 1850000, ["C3", "C6"], [[("C3", "C6")]]),
    # relationship: yes/no, including the similar-name traps
    ("Q15", "yes_no", "Is Kestrel Alloy Castings Limited a related party of Larkspur Castings?",
     "yes_no", "no", ["C9", "C1"], [[("C9", "C1")]]),
    ("Q16", "yes_no", "Is Larkspur Forgings Limited a subsidiary of Larkspur Castings?",
     "yes_no", "no", ["C10", "C1"], [[("C10", "C1")]]),
    ("Q17", "yes_no", "Is Arun Mehta a director of Larkspur Castings?",
     "yes_no", "no", ["P6"], []),
    ("Q18", "yes_no", "Is Kavita Bose a director of Bluefin Traders?",
     "yes_no", "yes", ["P4", "C6"], [[("P4", "C6")]]),
    # multi-hop: "how is X connected to Y", gold = the intermediate entities
    ("Q19", "multi_hop", "How is Kavita Bose connected to Larkspur Castings?",
     "entity", ["C3"], ["P4", "C3", "C1"], [[("P4", "C3"), ("C1", "C3")]]),
    ("Q20", "multi_hop", "How is Arun Mehta connected to Larkspur Castings?",
     "entity", ["C7"], ["P6", "C7", "C1"], [[("P6", "C7"), ("C7", "C1")]]),
    ("Q21", "multi_hop", "How is Tomas Fernandes connected to Larkspur Castings?",
     "entity_any", ["C4", "C5"], ["P8", "C1"],
     [[("P8", "C4"), ("C1", "C4")], [("P8", "C5"), ("C1", "C5")]]),
    ("Q22", "multi_hop", "How is Vikram Sethi connected to Harbourline Logistics?",
     "entity", ["C2", "C1"], ["P9", "C2", "C1", "C3"], [[("P9", "C2"), ("C2", "C1"), ("C1", "C3")]]),
    ("Q23", "multi_hop", "How is Kavita Bose connected to Tern Capital?",
     "entity", ["C3", "C1"], ["P4", "C3", "C1", "C2"], [[("P4", "C3"), ("C1", "C3"), ("C2", "C1")]]),
    ("Q24", "multi_hop", "How is Bluefin Traders connected to Larkspur Foundry Services?",
     "entity", ["C3", "C1"], ["C6", "C3", "C1", "C4"], [[("C3", "C6"), ("C1", "C3"), ("C1", "C4")]]),
    ("Q25", "multi_hop", "How is Patel Rao & Associates connected to Tern Capital?",
     "entity", ["C3", "C1"], ["C8", "C3", "C1", "C2"], [[("C8", "C3"), ("C1", "C3"), ("C2", "C1")]]),
    ("Q26", "multi_hop", "How is Tomas Fernandes connected to Harbourline Logistics?",
     "entity", ["C1"], ["P8", "C1", "C3"],
     [[("P8", "C4"), ("C1", "C4"), ("C1", "C3")], [("P8", "C5"), ("C1", "C5"), ("C1", "C3")]]),
    # cross-document: the facts needed are never stated together in one document
    # (checked by the generator: min_documents >= 2)
    ("Q27", "cross_doc", "Which person is a director of both Larkspur Castings and Harbourline Logistics?",
     "entity", ["P1"], ["P1", "C1", "C3"],
     [[("P1", "C1", "MANAGING_DIRECTOR_OF"), ("P1", "C3", "DIRECTOR_OF")]]),
    ("Q28", "cross_doc", "Which director of Larkspur Castings also sits on the board of Larkspur Foundry "
                         "Services?",
     "entity", ["P3"], ["P3", "C1", "C4"],
     [[("P3", "C1", "WHOLE_TIME_DIRECTOR_OF"), ("P3", "C4", "CHAIRPERSON_OF")]]),
    ("Q29", "cross_doc", "Which designated partner of Kestrel Alloys LLP is also a director of Harbourline "
                         "Logistics?",
     "entity", ["P1"], ["P1", "C5", "C3"],
     [[("P1", "C5", "DESIGNATED_PARTNER_OF"), ("P1", "C3", "DIRECTOR_OF")]]),
    ("Q30", "cross_doc", "Which audit partner leads the statutory audit of the company whose board meeting "
                         "Rohan Pillai chaired on 22 May 2025?",
     "entity", ["P6"], ["P6", "C7", "C4", "P3"],
     [[("P3", "C4", "CHAIRPERSON_OF"), ("C7", "C4", "AUDITOR_OF"), ("P6", "C7", "PARTNER_OF")]]),
    ("Q31", "cross_doc", "What is the DIN of the director who chaired the Larkspur Castings board meeting "
                         "of 18 July 2025?",
     "string", "07410101", ["P1", "C1"], [[("P1", "C1", "CHAIRPERSON_OF")]]),
    ("Q32", "cross_doc", "How many shares of Larkspur Castings are held by its Managing Director?",
     "number", 1500000, ["P1", "C1"], [[("P1", "C1", "MANAGING_DIRECTOR_OF")]]),
    ("Q33", "cross_doc", "How much alloy scrap, in INR per year, does the subsidiary chaired by Larkspur "
                         "Castings' Chief Financial Officer buy from Kestrel Alloys?",
     "number", 2100000, ["P3", "C4", "C5"],
     [[("P3", "C1", "CFO_OF"), ("P3", "C4", "CHAIRPERSON_OF"), ("C4", "C5", "PURCHASED_FROM")]]),
    # list: entity-set answers, scored by F1
    ("Q34", "list", "List all directors of Larkspur Castings Private Limited.",
     "entity_list", ["P1", "P2", "P3", "P7"], ["P1", "P2", "P3", "P7", "C1"],
     [[("P1", "C1"), ("P2", "C1"), ("P3", "C1"), ("P7", "C1")]]),
    ("Q35", "list", "List the subsidiaries of Larkspur Castings.",
     "entity_list", ["C3", "C4"], ["C3", "C4", "C1"], [[("C1", "C3"), ("C1", "C4")]]),
    ("Q36", "list", "List the directors of Harbourline Logistics.",
     "entity_list", ["P1", "P4"], ["P1", "P4", "C3"], [[("P1", "C3"), ("P4", "C3")]]),
    ("Q37", "list", "Which companies does Sen & Varma LLP audit?",
     "entity_list", ["C1", "C4"], ["C7", "C1", "C4"], [[("C7", "C1"), ("C7", "C4")]]),
    ("Q38", "list", "List the designated partners of Kestrel Alloys LLP.",
     "entity_list", ["P1", "P8"], ["P1", "P8", "C5"], [[("P1", "C5"), ("P8", "C5")]]),
    ("Q39", "list", "List the directors of Tern Capital.",
     "entity_list", ["P7", "P9"], ["P7", "P9", "C2"], [[("P7", "C2"), ("P9", "C2")]]),
]
TYPE_ORDER = ["lookup", "attribute", "yes_no", "multi_hop", "cross_doc", "list"]


# --------------------------------------------------------------------------
# Rendering helpers
# --------------------------------------------------------------------------
def render(text, doc, page, occurrences, entity_ids):
    """Replace [[ID|surface]] with surface and record the occurrence."""
    if text is None or not isinstance(text, str):
        return text

    def repl(match):
        eid, surface = match.group(1), match.group(2)
        if eid not in entity_ids:
            raise ValueError(f"{doc}: unknown entity {eid}")
        occurrences.append({"document": doc, "page": page, "entity_id": eid, "surface": surface})
        return surface

    return MARKUP.sub(repl, text)


def _pdf_lines(blocks, width=92):
    lines = []
    for kind, text in blocks:
        if kind == "h":
            lines.append(("h", text))
        else:
            for piece in textwrap.wrap(text, width=width) or [""]:
                lines.append(("p", piece))
        lines.append(("gap", ""))
    return lines


def write_pdf(path, pages, title):
    import fitz  # PyMuPDF

    doc = fitz.open()
    for blocks in pages:
        page = doc.new_page(width=595, height=842)  # A4
        y = 56
        for kind, text in _pdf_lines(blocks):
            if kind == "gap":
                y += 5
                continue
            size = 12 if kind == "h" else 10
            font = "helv" if kind == "p" else "hebo"
            page.insert_text((50, y), text, fontsize=size, fontname=font)
            y += size + 5
        if y > 800:
            raise ValueError(f"{path.name}: page overflow")
    stamp = "D:20250101000000"
    doc.set_metadata({"title": title, "author": "Synthetic benchmark", "creator": "", "producer": "",
                      "creationDate": stamp, "modDate": stamp})
    doc.save(str(path), garbage=4, deflate=True, no_new_id=True)
    doc.close()


def write_docx(path, pages, title):
    from docx import Document

    document = Document()
    for blocks in pages:
        for kind, text in blocks:
            if kind == "h":
                document.add_heading(text, level=1)
            else:
                document.add_paragraph(text)
    props = document.core_properties
    props.author = "Synthetic benchmark"
    props.last_modified_by = "Synthetic benchmark"
    props.title = title
    props.comments = NOTE
    props.created = props.modified = props.last_printed = FIXED_TIME
    props.revision = 1
    document.save(str(path))
    _freeze_zip(path)


def write_xlsx(path, sheets, title):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet["sheet"])
        for row in sheet["rows"]:
            ws.append(row)
    wb.properties.creator = "Synthetic benchmark"
    wb.properties.lastModifiedBy = "Synthetic benchmark"
    wb.properties.title = title
    wb.properties.created = wb.properties.modified = FIXED_TIME
    wb.save(str(path))
    _freeze_zip(path)


def _freeze_zip(path):
    """Rewrite a zip container with fixed timestamps so the output is byte-identical."""
    with zipfile.ZipFile(path) as src:
        items = [(info, src.read(info.filename)) for info in src.infolist()]
    xlsx_time = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, data in items:
            if info.filename == "docProps/core.xml":
                # openpyxl writes the save time into dcterms:modified
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + xlsx_time.encode() + rb"\g<2>", data)
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED_TIME.timetuple()[:6])
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o600 << 16
            dst.writestr(fixed, data)
    path.write_bytes(buf.getvalue())


# --------------------------------------------------------------------------
# Ground truth
# --------------------------------------------------------------------------
def _type_prefix(entity_type):
    return {"Person": "PERSON", "Company": "COMPANY"}.get(entity_type, "ORG")


def build_ground_truth(occurrences, doc_texts):
    entities = {e["id"]: e for e in ENTITIES}

    # Every surface used in the text must be a declared alias or ambiguous form.
    for occ in occurrences:
        ent = entities[occ["entity_id"]]
        if occ["surface"] not in ent["aliases"] + ent.get("ambiguous", []):
            raise ValueError(f"{occ['document']}: '{occ['surface']}' is not declared for {ent['id']}")

    mentions, mention_relations = [], []
    for doc in DOCUMENTS:
        doc_occ = [o for o in occurrences if o["document"] == doc["name"]]
        counters, by_entity = {}, {}
        for eid, name, attrs in doc["mentions"]:
            ent = entities[eid]
            if not any(o["entity_id"] == eid and o["surface"] == name for o in doc_occ):
                raise ValueError(f"{doc['name']}: mention '{name}' for {eid} does not occur in the text")
            prefix = _type_prefix(ent["type"])
            counters[prefix] = counters.get(prefix, 0) + 1
            pages = sorted({o["page"] for o in doc_occ if o["entity_id"] == eid})
            mention = {
                "mention_id": f"{doc['name']}:{prefix}_{counters[prefix]:03d}",
                "document": doc["name"],
                "extractor_id": f"{prefix}_{counters[prefix]:03d}",
                "entity_id": eid,
                "type": ent["type"],
                "attributes": {"name": name, **attrs, "page_numbers": pages},
            }
            mentions.append(mention)
            by_entity[eid] = mention["mention_id"]
        missing = {o["entity_id"] for o in doc_occ} - set(by_entity)
        if missing:
            raise ValueError(f"{doc['name']}: entities {sorted(missing)} occur but have no mention")
        for src, rel, dst, attrs in doc["relations"]:
            mention_relations.append({"document": doc["name"], "from": by_entity[src], "to": by_entity[dst],
                                      "type": rel, "attributes": attrs})

    facts = {}
    for doc in DOCUMENTS:
        for src, rel, dst, attrs in doc["relations"]:
            key = (src, rel, dst)
            fact = facts.setdefault(key, {"from": src, "type": rel, "to": dst, "attributes": {},
                                          "documents": []})
            fact["attributes"].update(attrs)
            if doc["name"] not in fact["documents"]:
                fact["documents"].append(doc["name"])
    relations = sorted(facts.values(), key=lambda f: (f["from"], f["type"], f["to"]))
    def fact_docs(pair):
        """Documents stating a relation between the pair (of the given type, if any)."""
        a, b = pair[0], pair[1]
        rel_type = pair[2] if len(pair) > 2 else None
        return {d for f in relations if {f["from"], f["to"]} == {a, b}
                and (rel_type is None or f["type"] == rel_type) for d in f["documents"]}

    def value_docs(answer_type, answer):
        """Documents whose text states the gold number or identifier."""
        from benchmarks.scoring import normalize_text, parse_numbers

        if answer_type == "number":
            return {d for d, t in doc_texts.items() if any(abs(v - answer) < 1e-6 for v in parse_numbers(t))}
        if answer_type == "string":
            return {d for d, t in doc_texts.items() if str(answer) in normalize_text(t).split()}
        return None

    def min_cover(requirements):
        """Smallest set of documents that meets every requirement (each a set of documents)."""
        names = sorted(doc_texts)
        for size in range(1, len(names) + 1):
            for combo in itertools.combinations(names, size):
                if all(req & set(combo) for req in requirements):
                    return list(combo)
        raise ValueError("requirements cannot be met")

    questions = []
    rank = {}
    for qid, qtype, text, answer_type, answer, evidence, paths in QUESTIONS:
        rank[qtype] = rank.get(qtype, 0) + 1
        gold_entities = answer if answer_type.startswith("entity") else []
        for eid in list(gold_entities) + list(evidence):
            if eid not in entities:
                raise ValueError(f"{qid}: unknown entity {eid}")
        covers = []
        for alt in paths or [[]]:
            requirements = []
            for pair in alt:
                docs = fact_docs(pair)
                if not docs:
                    raise ValueError(f"{qid}: no relation fact supports {pair}")
                requirements.append(docs)
            vdocs = value_docs(answer_type, answer)
            if vdocs is not None:
                if not vdocs:
                    raise ValueError(f"{qid}: gold value {answer!r} does not occur in the corpus")
                requirements.append(vdocs)
            if requirements:
                covers.append(min_cover(requirements))
        needed = min(covers, key=len) if covers else []
        if qtype == "cross_doc" and len(needed) < 2:
            raise ValueError(f"{qid}: a cross-document question is answerable from {needed}")
        questions.append({
            "id": qid, "type": qtype, "question": text, "answer_type": answer_type,
            "answer": answer, "evidence_entities": evidence,
            "evidence_paths": [[list(p) for p in alt] for alt in paths],
            # Fewest documents that together state every fact the question needs, and one such set.
            "min_documents": len(needed) if needed else None,
            "documents_needed": needed,
            # Round-robin order over question types: any prefix of the run order is stratified.
            "priority": rank[qtype] * 10 + TYPE_ORDER.index(qtype),
        })

    entity_list = [{"id": e["id"], "type": e["type"], "canonical_name": e["name"],
                    "aliases": e["aliases"], "ambiguous_surfaces": e.get("ambiguous", [])}
                   for e in ENTITIES]
    return {
        "entities.json": {"entities": entity_list},
        "mentions.json": {"mentions": mentions, "relations": mention_relations},
        "relations.json": {"relations": relations},
        "questions.json": {"questions": questions},
        "occurrences.json": {"occurrences": occurrences},
    }


def build(out_dir: Path = BENCH_DIR):
    corpus_dir = out_dir / "corpus"
    truth_dir = out_dir / "ground_truth"
    corpus_dir.mkdir(parents=True, exist_ok=True)
    truth_dir.mkdir(parents=True, exist_ok=True)
    entity_ids = {e["id"] for e in ENTITIES}

    occurrences = []
    documents = []
    doc_texts = {}
    for doc in DOCUMENTS:
        name, fmt = doc["name"], doc["format"]
        if fmt == "xlsx":
            sheets = [{"sheet": s["sheet"],
                       "rows": [[render(c, name, i, occurrences, entity_ids) for c in row]
                                for row in s["rows"]]}
                      for i, s in enumerate(doc["pages"], 1)]
            write_xlsx(corpus_dir / f"{name}.xlsx", sheets, doc["title"])
            doc_texts[name] = "\n".join(" | ".join("" if c is None else str(c) for c in row)
                                         for sheet in sheets for row in sheet["rows"])
        else:
            pages = [[(kind, render(text, name, i, occurrences, entity_ids)) for kind, text in blocks]
                     for i, blocks in enumerate(doc["pages"], 1)]
            writer = write_pdf if fmt == "pdf" else write_docx
            writer(corpus_dir / f"{name}.{fmt}", pages, doc["title"])
            doc_texts[name] = "\n".join(text for blocks in pages for _, text in blocks)
        documents.append({"name": name, "file": f"{name}.{fmt}", "format": fmt,
                          "pages": len(doc["pages"])})

    truth = build_ground_truth(occurrences, doc_texts)
    truth["documents.json"] = {"documents": documents}
    for filename, data in truth.items():
        (truth_dir / filename).write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n",
                                          encoding="utf-8")
    return corpus_dir, truth_dir


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", type=Path, default=BENCH_DIR,
                        help="directory that receives corpus/ and ground_truth/ (default: benchmarks/)")
    args = parser.parse_args()
    corpus_dir, truth_dir = build(args.out)
    n_mentions = len(json.loads((truth_dir / "mentions.json").read_text())["mentions"])
    print(f"Wrote {len(DOCUMENTS)} documents to {corpus_dir}")
    print(f"Wrote ground truth to {truth_dir}: {len(ENTITIES)} entities, {n_mentions} mentions, "
          f"{len(QUESTIONS)} questions")


if __name__ == "__main__":
    main()
